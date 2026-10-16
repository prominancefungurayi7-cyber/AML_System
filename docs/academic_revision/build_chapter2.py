from pathlib import Path
import ast
import re
import json
from docx import Document
from docx.shared import Pt
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent
# Reuse formatting helpers only; never execute the Chapter 1 document builder.
tree = ast.parse((OUT / 'build_revision.py').read_text(encoding='utf-8'))
helpers = ast.Module(body=[n for n in tree.body if isinstance(n, (ast.Import, ast.ImportFrom, ast.FunctionDef))], type_ignores=[])
exec(compile(helpers, 'chapter1_format_helpers', 'exec'), globals())

refs = {
'gsma': 'GSMA, The State of the Industry Report on Mobile Money 2025, 2025. [Online]. Available: https://www.gsma.com/sotir/wp-content/uploads/2025/04/The-State-of-the-Industry-Report-2025_English.pdf. Accessed: Oct. 14, 2026.',
'findex': 'A. Demirgüç-Kunt, L. Klapper, D. Singer, and S. Ansar, The Global Findex Database 2021: Financial Inclusion, Digital Payments, and Resilience in the Age of COVID-19. Washington, DC, USA: World Bank, 2022, doi: 10.1596/978-1-4648-1897-4.',
'whisker': 'J. Whisker and M. E. Lokanan, “Anti-money laundering and counter-terrorist financing threats posed by mobile money,” Journal of Money Laundering Control, vol. 22, no. 1, pp. 158–172, 2019, doi: 10.1108/JMLC-10-2017-0061.',
'interpol': 'INTERPOL, Mobile Money and Organized Crime in Africa, June 2020. [Online]. Available: https://www.interpol.int/en/content/download/15457/file/2020%2007%2015%20PUBLIC%20VERSION%20-%20Strategic%20Analysis%20Report%20-Mobile%20Money%20in%20Africa%202020.pdf. Accessed: Oct. 14, 2026.',
'jensen': 'R. I. T. Jensen and A. Iosifidis, “Fighting money laundering with statistics and machine learning,” IEEE Access, vol. 11, pp. 8889–8903, 2023, doi: 10.1109/ACCESS.2023.3239549.',
'altman': 'E. Altman, J. Blanuša, L. von Niederhäusern, B. Egressy, A. Anghel, and K. Atasu, “Realistic synthetic financial transactions for Anti-Money Laundering models,” Advances in Neural Information Processing Systems, vol. 36, 2023. [Online]. Available: https://arxiv.org/abs/2306.16424. Accessed: Oct. 14, 2026.',
'weber': 'M. Weber et al., “Anti-Money Laundering in Bitcoin: Experimenting with Graph Convolutional Networks for Financial Forensics,” KDD Workshop on Anomaly Detection in Finance, 2019. [Online]. Available: https://arxiv.org/abs/1908.02591. Accessed: Oct. 14, 2026.',
'sowon': 'K. Sowon, E. Luhanga, L. F. Cranor, G. Fanti, C. Tucker, and A. Gueye, “The role of user-agent interactions on mobile money practices in Kenya and Tanzania,” in Proc. IEEE Symposium on Security and Privacy, 2024, pp. 752–769, doi: 10.1109/SP54263.2024.00184.',
'mogaji': 'E. Mogaji and N. P. Nguyen, “The dark side of mobile money: Perspectives from an emerging economy,” Technological Forecasting and Social Change, vol. 185, Art. no. 122045, 2022, doi: 10.1016/j.techfore.2022.122045.',
'chitimira2023': 'H. Chitimira and E. Torerai, “Policy implications and mobile money regulatory approaches to promote financial inclusion of the poor in Zimbabwe after the COVID-19 pandemic,” De Jure, vol. 56, no. 1, 2023, doi: 10.17159/2225-7160/2023/v56a17.',
'fatf': 'Financial Action Task Force, Opportunities and Challenges of New Technologies for AML/CFT, Paris, France, 2021. [Online]. Available: https://www.fatf-gafi.org/content/dam/fatf-gafi/guidance/Opportunities-Challenges-of-New-Technologies-for-AML-CFT.pdf. Accessed: Oct. 14, 2026.',
'jullum': 'M. Jullum, A. Løland, R. B. Huseby, G. Ånonsen, and J. Lorentzen, “Detecting money laundering transactions with machine learning,” Journal of Money Laundering Control, vol. 23, no. 1, pp. 173–186, 2020, doi: 10.1108/JMLC-07-2019-0055.',
'tertychnyi': 'P. Tertychnyi, M. Godgildieva, M. Dumas, and M. Ollikainen, “Time-aware and interpretable predictive monitoring system for Anti-Money Laundering,” Machine Learning with Applications, vol. 8, Art. no. 100306, 2022, doi: 10.1016/j.mlwa.2022.100306.',
'kapoor': 'S. Kapoor and A. Narayanan, “Leakage and the reproducibility crisis in machine-learning-based science,” Patterns, vol. 4, no. 9, Art. no. 100804, 2023, doi: 10.1016/j.patter.2023.100804.',
'mazumder': 'P. T. Mazumder, “Explainable and fair anti-money laundering models using a reproducible SHAP framework for financial institutions,” Discover Artificial Intelligence, vol. 6, Art. no. 262, 2026, doi: 10.1007/s44163-026-00944-7.',
'esaamlg': 'Eastern and Southern Africa Anti-Money Laundering Group, Zimbabwe: 10th Enhanced Follow-Up Report and 4th Technical Compliance Re-Rating, Apr. 2024. [Online]. Available: https://www.esaamlg.org/reports/Zimbabwe_FUR-April%202024.pdf. Accessed: Oct. 14, 2026.',
'rbz': 'Reserve Bank of Zimbabwe, Zimbabwe National Financial Inclusion Strategy II, 2022–2026, 2022. [Online]. Available: https://www.rbz.co.zw/documents/BLSS/FinancialInclusion/Zimbabwe_National_Financial_Inclusion_Strategy_II_2022-2026.pdf. Accessed: Oct. 14, 2026.',
'chitimira2021': 'H. Chitimira and E. Torerai, “The nexus between mobile money regulation, innovative technology and the promotion of financial inclusion in Zimbabwe,” Potchefstroom Electronic Law Journal, vol. 24, 2021, doi: 10.17159/1727-3781/2021/v24i0a10739.',
'ndlovu': 'B. Ndlovu, F. J. Kiwa, M. Muduva, C. T. Chipfumbu, S. Marambi, and A. Maphosa, “Developing a graph-based machine learning model for identifying money laundering networks associated with sanctioned entities in a bank in Zimbabwe,” Indonesian Journal of Innovation and Applied Sciences, vol. 6, no. 1, pp. 35–59, 2026, doi: 10.47540/ijias.v6i1.2306.'
}
order=[]
def cite(text):
    def sub(m):
        key=m.group(1)
        assert key in refs, key
        if key not in order: order.append(key)
        return '['+str(order.index(key)+1)+']'
    return re.sub(r'\{cite:([a-z0-9]+)\}', sub, text)

d=base('Chapter 2: Literature Review')
d.core_properties.subject='Revised literature review responding to 15 supervisor comments'
d.sections[0].footer.paragraphs[0].clear()
for _ in range(5):d.add_paragraph()
for text in ['Mobile-Money Anti-Money Laundering Decision-Support System','Chapter 2: Literature Review','Prominance Fungurayi: R251946T','Ashton Kahlari: R252989W']:
    par=p(d,text);par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if text.startswith(('Mobile','Chapter')):
        for run in par.runs:run.bold=True;run.font.size=Pt(14)
s=d.add_section(WD_SECTION_START.NEW_PAGE);s.footer.is_linked_to_previous=False
num=OxmlElement('w:pgNumType');num.set(qn('w:start'),'1');s._sectPr.append(num)
par=s.footer.paragraphs[0];par.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');par._p.append(fld)
def para(text):return p(d,cite(text))
def heading(text,level=3):h(d,text,level)

heading('CHAPTER 2: LITERATURE REVIEW',1)
heading('2.1 Introduction',2)
para('This chapter critically reviewed literature relevant to mobile-money Anti-Money Laundering (AML) decision support. The review followed a funnel structure, beginning with international financial monitoring and analytical methods, then considering African mobile-money evidence and finally Zimbabwe. It examined structuring, wallet relationships, agent activity, rules and learning methods, data quality, evaluation, explainability and timely investigation. Studies were compared by research question, method, data, findings, limitations and relevance to the proposed application. The purpose was to establish a defensible research gap and motivate the proposed work, rather than report its implementation or results.')
para('For this review, mobile money referred to value held and transferred through mobile-wallet services, including agent-mediated cash conversion. A behavioural feature was a numerical description of transaction history or relationships. A suspicious-pattern classification indicated a case for assessment, not a determination of criminal conduct. An operational alert was a record created by a monitoring control; its existence did not make it a verified training label.')
heading('2.2 Analysis of Related Work',2)
para('The review combined peer-reviewed empirical and methodological studies with institutional publications used for industry and regulatory context. The source set was a focused narrative review, not an exhaustive systematic review. Consequently, the gaps identified below concern the evidence examined and do not imply that equivalent proprietary monitoring systems do not exist.')

heading('2.2.1 Mobile Money and the Digital Financial Environment')
para('The GSMA’s 2025 industry report documented the expansion of mobile-money services and the central role of Sub-Saharan Africa {cite:gsma}. Demirgüç-Kunt et al. (2022), using the Global Findex household-survey evidence, described financial access and the growing use of digital payments {cite:findex}. These sources established the significance of the service environment. However, account ownership and self-reported usage were not labelled transaction histories and could not validate an AML classifier.')
para('Whisker and Lokanan (2019) reviewed mobile-money characteristics and analysed themes relevant to AML and counter-terrorist-financing controls {cite:whisker}. Their findings supported attention to the vulnerabilities of payment channels and their controls. The study examined risk mechanisms rather than measuring the predictive performance of a transaction detector. It therefore provided a domain rationale, while leaving the effectiveness of particular monitoring methods to empirical evaluation.')
para('A banking detector could offer useful methods without being directly transferable to mobile money. The latter’s cash-in/cash-out interface introduced agent relationships alongside wallet transfers; a customer’s activity could span both. This review therefore treated the monitored unit as a transaction embedded in a history and service network. That analytical choice required evaluation rather than an assumption that all frequent transfers or agent interactions were suspicious.')

heading('2.2.2 Money-Laundering Risk in Mobile-Money Environments')
para('INTERPOL’s 2020 strategic analysis documented criminal exploitation of African mobile-money services, including laundering and fraud, and identified weaknesses in identity checks and investigative capacity {cite:interpol}. This was stronger evidence of observed misuse than a statement that abuse was merely possible. Nevertheless, a strategic threat assessment did not provide a representative estimate of the proportion of illicit transactions or a labelled benchmark for classification.')
para('The risk literature supported examining how funds moved and how controls could be bypassed. It did not justify treating every unusual payment as illicit. Fraud, laundering, identity misuse and customer inconvenience also represented different phenomena: evidence of one could not automatically establish another. The implication for this study was to preserve the distinction between a suspicious-pattern indicator, a rule breach and an investigator’s conclusion {cite:whisker}.')

heading('2.2.3 Structuring Behaviour')
para('Structuring concerned the distribution of value across multiple transactions in ways that could obscure the overall movement. Jensen and Iosifidis (2023) reviewed statistical and machine-learning AML research and distinguished client-risk profiling from suspicious-behaviour flagging {cite:jensen}. Their banking-focused review identified challenges involving data availability and the interpretation of risk indicators. It supported examining behaviour across observations, but did not validate a mobile-money structuring detector.')
para('Altman et al. (2023) provided synthetic AML datasets in which specified laundering patterns could be evaluated under controlled conditions {cite:altman}. Such experiments enabled investigation of repeated or distributed transfers, but their conclusions remained tied to the generated scenarios. For the present study, this motivated testing whether historical calculations represented the intended patterns, separately from testing whether a classifier detected them. Frequent small payments could be legitimate; an amount or frequency threshold alone could not establish intent.')

heading('2.2.4 Wallet and Transaction-Network Behaviour')
para('Weber et al. (2019) compared conventional classifiers and graph convolutional networks on the Elliptic Bitcoin transaction dataset {cite:weber}. Random Forest performed strongly in their experiments, demonstrating that a graph-learning architecture was not automatically superior. The data concerned cryptocurrency transactions with partially available labels, so the findings did not directly establish performance for mobile-wallet cash flows or agents.')
para('The distinction was between using relational information and learning directly over an entire graph. A full graph model could represent connected neighbourhoods, but required decisions about graph boundaries, time, labels and evaluation. Network-derived tabular features summarised selected relationships into quantities that could be checked against transaction records. The proposed study adopted the latter as a bounded, auditable approach compatible with its local prototype and 12-week scope. This was a scope and verification decision, not a claim of greater predictive accuracy. Summaries could discard multi-hop structures and other relationships that a full graph model might retain; this remained a limitation.')

heading('2.2.5 Mobile-Money Agent Behaviour')
para('Sowon et al. (2024) conducted 72 structured interviews with mobile-money users in Kenya and Tanzania. They identified workarounds involving agents, identification and transaction execution {cite:sowon}. Their findings showed that observed activity reflected social and operational arrangements as well as formal service procedures. The qualitative sample could explain practices but could not estimate laundering prevalence or train a reliable AML classifier. It informed the need to interpret repeated customer-agent interactions in context.')
para('Mogaji and Nguyen (2022) analysed interviews with financially vulnerable customers, agents and financial-technology developers. They identified stakeholder relationships involving fraud, privacy, security and support needs {cite:mogaji}. The study added evidence that agent-related risk extended beyond the amount of a single payment. Its thematic findings did not supply transaction-level ground truth or demonstrate that high-volume agents were criminal. Compared with that qualitative work, the proposed study investigated whether explicitly defined transaction-history summaries could support a testable monitoring process.')
para('Zimbabwean agent regulation also featured in Chitimira and Torerai’s (2023) legal analysis of mobile-money policy {cite:chitimira2023}, examined further in Section 2.2.14. Across these sources, agent practices were documented more clearly than the predictive validity of agent-history measurements. That specific gap motivated including agent behaviour alongside wallet and transaction histories. It did not justify assuming that concentration, repeated customers or growth in activity was sufficient evidence of laundering.')

heading('2.2.6 Conventional AML Monitoring, Labels and Operational Rules')
para('FATF (2021) examined opportunities for new AML/CFT technologies while emphasising implementation challenges involving data and governance {cite:fatf}. Rules remained useful for expressing known conditions, while learned models could combine information differently. Their coexistence required clear attribution of why an alert existed. A model score and an independently triggered rule were different outputs and could disagree without either being silently rewritten.')
para('Supervised learning required reference labels for training and evaluation. In an operational study, such labels might reflect recorded investigative outcomes, with their associated uncertainty. In a synthetic experiment, labels described the scenarios assigned by the generator {cite:altman}. These were experimental reference labels, not verified real-world criminality. They were distinct from application rules evaluated when a transaction was processed. Training a model on rule-derived labels was possible as a rule-reproduction task, but agreement would not establish independent laundering detection. The distinction was necessary to state clearly what the model learned and what its evaluation could establish.')
d.add_paragraph('Table 2.1: Distinction between labels and monitoring outputs','Caption')
table(d,['Item','Purpose and interpretation'],[
('Reference label','Defines the target for training or evaluation; its source and uncertainty must be documented.'),
('Model prediction','An estimated class or score produced from permitted inputs; evaluated against reference labels.'),
('Rule alert','A recorded operational condition triggered by a rule; not automatically a training label or a model explanation.')],[1.6,4.7])

heading('2.2.7 Machine Learning for AML Detection')
para('Jullum et al. (2020) trained a supervised model using DNB banking transactions, investigated alerts and cases reported to authorities. Their XGBoost-based approach improved prioritisation under the study’s comparison with the bank’s monitoring process {cite:jullum}. An important finding concerned the value of including investigated but unreported alerts. These outcomes reflected institutional reporting decisions, not universal certainty about criminality. The work was relevant to investigation support, but its banking data and institutional workflow differed from a synthetic mobile-money prototype.')
para('Tertychnyi et al. (2022) combined history-based prediction, an alert policy and interpretability in a monitoring system evaluated with institutional data and subject-matter-expert feedback {cite:tertychnyi}. Their work considered timely, non-redundant alerts as part of usefulness. This broadened evaluation beyond selecting a classifier. Its findings remained specific to the investigated setting and did not remove the need to assess local mobile-money data and workflows.')
para('Supervised, unsupervised and graph-based approaches addressed different information conditions. A supervised classifier required a defined target; an anomaly score identified unusual behaviour without necessarily matching that target; a graph method depended on the available relationships. The present study’s methodological implication was to define the task and compare candidates under one evaluation protocol. Detailed candidate settings, selection decisions and achieved results were reserved for Chapters 3 and 4.')

heading('2.2.8 Temporal Behaviour and Data Leakage')
para('Kapoor and Narayanan (2023) examined leakage across machine-learning research and showed how information unavailable in the intended prediction setting could inflate apparent performance {cite:kapoor}. Their analysis was broader than AML, but its methodological relevance was direct: a retrospective dataset could contain future outcomes, repeated entities or preprocessing information that would be unavailable to a live decision. Their work did not provide an AML-specific correction that could be assumed sufficient in every dataset.')
para('For transaction monitoring, two questions therefore had to remain separate: whether each feature used information available at its decision time, and whether the evaluation partition supported the intended generalisation claim. Chronological feature construction alone did not prove independence of evaluation data, while a held-out partition did not repair features containing future information. The review established these requirements; the project’s ordering, partitioning and verification mechanisms belonged in Chapter 3.')

heading('2.2.9 Synthetic Data for AML Research')
para('Altman et al. (2023) developed a configurable financial-transaction generator and benchmark datasets to support reproducible AML experiments where operational records were difficult to share {cite:altman}. Their contribution enabled comparisons with known simulated scenarios. However, models could learn regularities of a generator, and scenario labels did not reproduce the uncertainty of real investigations. The relevant benefit was controlled experimentation; the limitation was restricted external validity.')
para('This review therefore treated synthetic evidence as evidence about the defined simulation. A convincing evaluation needed to state what was simulated, how labels were assigned, which observations were used for model development and what remained independently assessed. Synthetic success could not establish effectiveness on an operator network. The project-specific dataset size, composition and generation procedure were omitted here and reserved for the methodology chapter.')

heading('2.2.10 Class Imbalance and AML Evaluation')
para('The AML literature identified sparse labels and uneven class frequencies as obstacles to interpretation {cite:jensen}. Accuracy alone could conceal missed suspicious cases. For example, in a hypothetical set of 1,000 observations containing ten positive labels, predicting every observation as negative would yield 99% accuracy and zero positive-class recall. This arithmetic example illustrated a metric limitation; it was not an estimate of real laundering prevalence.')
para('Precision described the proportion of positive predictions matching the reference class, while recall described the proportion of reference positives detected. False-positive rate measured false alerts among reference negatives; it was different from the proportion of alerts that were false. Confusion matrices and class-sensitive measures were therefore needed alongside aggregate scores. Comparisons also required the same task, label meaning and evaluation data. A scenario-generation ratio, an alert ratio and classification accuracy measured different things and should not be substituted for one another.')

heading('2.2.11 Explainability and Human Investigation')
para('Tertychnyi et al. (2022) included explanations and risk estimates in their investigator-oriented monitoring approach {cite:tertychnyi}. This provided academic support for assessing what information accompanied a prediction. Its use of expert feedback was particularly relevant, although it did not establish that any explanation format would be useful to every compliance team.')
para('Mazumder (2026) evaluated an explainability-oriented AML framework using a synthetic retail-banking dataset and SHAP-based explanations {cite:mazumder}. The work illustrated how local and global model explanations could be documented for scrutiny. Its synthetic setting limited claims about operational usefulness, and the generation of explanations was not itself evidence of improved investigator decisions. The present study’s comparison therefore concerned traceable decision support rather than an assumption that an explanation proved correctness.')
para('Three forms of information needed clear separation: a model attribution described a model’s output, a rule explanation described a triggered control, and a transaction-history summary described observed activity. A history summary should not be presented as a mathematical attribution unless that attribution was actually computed. Investigators could use these records to question a case, seek context and document a decision. The literature thus supported human review and provenance, without conferring authority on a model to determine criminal intent {cite:fatf}.')

heading('2.2.12 Timely AML Monitoring')
para('Time-aware monitoring research treated the timing and repetition of alerts as part of the monitoring problem {cite:tertychnyi}. A useful comparison therefore distinguished detection delay from computation time and from the time needed to deliver a notification. A fast estimator alone did not establish a fast investigation workflow.')
para('For the proposed study, “real-time” described the intention to make monitoring outcomes available during ongoing application use. A specific latency claim would require a defined start point, end point, workload and measurement. Training duration, inference duration and browser-update delay were different measurements. The literature motivated assessing the complete process, while the implementation sequence and measured timing results remained matters for Chapters 3 and 4.')

heading('2.2.13 African and Regional Perspective')
para('The regional evidence concerned both access and integrity. Findex survey evidence documented the role of mobile money in financial participation {cite:findex}, while the Kenya–Tanzania interview study and the stakeholder research identified practices and vulnerabilities that a transaction table alone could not fully explain {cite:sowon}, {cite:mogaji}. INTERPOL’s regional assessment supplied evidence of criminal exploitation, but it served a different purpose from either a representative household survey or an evaluated detector {cite:interpol}.')
para('These differences limited direct comparison: interview findings described experiences, industry reports described adoption and threat reports described risks. None supplied an interchangeable ground-truth dataset for a regional AML accuracy claim. The resulting regional gap was the limited evidence, within the sources reviewed, connecting observed user-agent practices to validated behavioural measurements and an evaluated investigation workflow. This motivated a controlled prototype while preserving the need for later operator-data and user evaluation.')
para('The Eastern and Southern Africa Anti-Money Laundering Group (ESAAMLG) provided a further institutional context. Its April 2024 Zimbabwe follow-up report assessed specified technical-compliance developments and explicitly did not assess improvements in effectiveness {cite:esaamlg}. The report could therefore support discussion of the assessment framework, but could not demonstrate that a particular monitoring algorithm worked or that the present prototype met regulatory requirements.')

heading('2.2.14 Zimbabwean Context')
para('Zimbabwe’s National Financial Inclusion Strategy II provided the policy context for access to and use of financial services {cite:rbz}. Chitimira and Torerai (2021) analysed mobile-money regulation and financial inclusion through legal and regulatory sources {cite:chitimira2021}. They identified tensions between technological change and the regulatory framework. Their analysis established local relevance but did not evaluate transaction-level prediction or investigator outcomes.')
para('Chitimira and Torerai (2023) compared regulatory approaches and discussed historical restrictions affecting agents and their implications for inclusion {cite:chitimira2023}. Their argument supported proportionate controls rather than treating access and integrity as separate concerns. It did not supply labelled agent histories or validate a behavioural classifier. These were findings about the period analysed, not an assertion that every regulatory condition remained unchanged at the time of this dissertation.')
para('Ndlovu et al. (2026) examined sanction-associated networks in a Zimbabwean banking setting using directed transaction graphs and a graph convolutional network compared with other classifiers {cite:ndlovu}. They reported an F1-score of 0.7345 and ROC-AUC of 0.8968. The article described banking transaction data, although its discussion of provenance and its slightly differing observation counts required caution when interpreting reproducibility. Its task concerned sanction-associated banking networks, not general mobile-money suspicious-pattern detection or agent-mediated cash activity. The reported scores were therefore not directly comparable with the present study’s synthetic classification results.')
para('The local evidence consequently included regulatory scholarship and an evaluated graph-learning study; it would be inaccurate to claim an absence of Zimbabwean AML research. The bounded gap was evidence for an integrated mobile-money prototype connecting transaction integrity, three behavioural dimensions, traceable monitoring outcomes and authorised investigation. The proposed study addressed that integration and demonstration task. It did not establish commercial novelty, replace operator controls or resolve the lack of independent real-world validation.')

heading('2.3 Gaps Identified',2)
para('Table 2.2 separates the limitations of the reviewed evidence from the proposed response. The response identifies what the study could evaluate within its scope; it does not claim that integration alone solved data scarcity, crime detection or operational deployment. The objective references correspond to the five SMART objectives in Chapter 1.')
d.add_paragraph('Table 2.2: Evidence-based gaps and bounded project responses','Caption')
gaps=[
('Transfer to mobile-money operations','Banking and Bitcoin studies assessed different tasks and data {cite:jullum}, {cite:weber}.','Evaluate transaction processing and behavioural analysis in a defined mobile-money setting (Objectives 1–3). No claim of universal transfer.'),
('Agent-behaviour validation','Interview and policy evidence did not validate agent-history classification {cite:sowon}, {cite:chitimira2023}.','Check agent-history calculations alongside structuring and wallet/network histories (Objective 2). Detection accuracy remains an empirical question.'),
('Reliable evaluation','Leakage and generator dependence could undermine interpretation {cite:kapoor}, {cite:altman}.','Document permitted information, labels and comparable evaluation (Objectives 2–3). External validation remains outside available evidence.'),
('Traceable investigator support','Prediction quality, explanation and alert utility were distinct evaluation concerns {cite:tertychnyi}, {cite:mazumder}.','Distinguish model and rule findings and test authorised case review (Objective 4). Do not claim proven productivity gains.'),
('Local workflow evidence','Zimbabwean regulation and graph research did not evaluate the complete mobile-money workflow specified here {cite:chitimira2021}, {cite:ndlovu}.','Test monitoring continuity, financial integrity and recovery alongside classification (Objectives 1 and 5). Operator deployment remains separate.')]
table(d,['Gap','Evidence and limitation','Proposed response'],[[cite(x) for x in row] for row in gaps],[1.3,2.2,2.8])

heading('2.4 Proposed Work',2)
para('The proposed work was a local mobile-money AML decision-support prototype combining transaction processing, historical behavioural analysis, machine-learning classification and authorised investigation. It considered structuring, wallet/network and agent behaviour, while keeping independent rule alerts distinguishable from model predictions. Its evaluation was tied to the five SMART objectives and the 12-week schedule established in Chapter 1.')
para('The intended contribution was a testable integration of these processes with traceable decisions and preserved financial records. It did not depend on claiming a new learning algorithm or proven real-world detection performance. Dataset construction, feature definitions, model comparison, implementation order and detailed recovery mechanisms were reserved for Chapter 3; measured outcomes were reserved for Chapter 4.')

heading('2.4.1 Feasibility Analysis')
para('Feasibility was considered at the scale of a local academic prototype. The assessment distinguished available resources from benefits and capabilities that still required measurement.')
heading('2.4.1.1 Technical Feasibility')
para('The existing local computer and the application’s Python, Flask and MySQL environment provided a practical basis for development. The software instruments were documented in Chapter 1. A tabular behavioural approach supported inspection of calculations without requiring a full graph-learning deployment. This established a reasonable implementation basis, not proof of concurrent-user capacity, complete security or operator-scale performance. Those claims required their own tests.')
heading('2.4.1.2 Economic Feasibility')
para('The incremental budget assumed reuse of the existing computer and local development tools. On that basis, additional hardware procurement and licence fees for the proposed open-source application stack were budgeted at US$0. This did not mean that the project had no cost: electricity, connectivity, development time and maintenance still had economic value. No verified expenditure record or supplier quotations were supplied for those items, so a total project cost and monetary return were not asserted.')
para('The anticipated benefit was a demonstrable investigation-support workflow and reusable research evidence. Reduced review effort was a possible future benefit, not a measured saving. Economic feasibility was therefore conditional on the existing resources and 12-week academic scope; a production business case would require separately costed infrastructure, integration, support and evaluation.')
heading('2.4.1.3 Social Feasibility')
para('Synthetic records reduced the need to expose confidential customer transactions during development. Human case review and explicit alert provenance supported scrutiny of outcomes that might otherwise be misunderstood. These design choices were consistent with the concerns raised in the literature, but they did not establish user acceptance, fairness across populations or absence of harm. Such conclusions would require stakeholder participation and suitable evaluation.')
heading('2.4.1.4 Operational Feasibility')
para('The proposed processes were compatible with a local workflow involving customer transactions and authorised compliance review. Operational feasibility depended on successful tests of transaction integrity, access control, alert handling and recovery. A failed inference could not be treated as a completed prediction or justify duplicating a payment. The existence of interface pages alone would not establish that these processes worked. Chapter 4 would report the relevant observed outcomes.')
heading('2.4.1.5 Feasibility Overview')
d.add_paragraph('Table 2.3: Feasibility assessment and evidence boundaries','Caption')
table(d,['Area','Assessment within the study'],[
('Technical','Development resources were available; capacity and reliability still required measurement.'),
('Economic','Reusing hardware and the open-source stack avoided additional procurement assumptions; total cost and savings were not verified.'),
('Social','Synthetic data and human review supported responsible design; stakeholder acceptance and fairness were not established.'),
('Operational','The workflow was suitable for controlled demonstration, subject to the acceptance tests in Chapter 1.')],[1.3,5.0])
para('The assessment supported bounded local development. It did not constitute approval for deployment on an operator network or evidence of regulatory compliance.')

heading('2.5 Summary',2)
para('This chapter compared international analytical research with African mobile-money evidence and Zimbabwean studies. The review supported attention to historical and relational behaviour, agent context, evaluation quality and investigator needs. It also distinguished reference labels, model predictions and operational rules. The identified gap concerned a demonstrable local integration of these processes, with explicit limits on what synthetic and institutional evidence could establish. The proposed work and feasibility assessment remained within that scope. Chapter 3 presents the methodology, architecture, data preparation, behavioural calculations, model-development procedure and integration design.')

heading('2.6 References',2)
for i,key in enumerate(order,1):
    par=p(d,f'[{i}] '+refs[key]);par.paragraph_format.keep_together=True
for par in d.paragraphs:
    if par.style.name=='Caption':par.paragraph_format.keep_with_next=True
    if par.style.name.startswith('Heading'):par.paragraph_format.keep_with_next=True
d.save(OUT/'Chapter_2_Corrected.docx')

g=base('Chapter 2: Supervisor Corrections Record')
g.core_properties.subject='Responses to Chapter 2 supervisor comments'
h(g,'CHAPTER 2: SUPERVISOR CORRECTIONS RECORD',1)
p(g,'The revision addresses all 15 embedded comments in Chapter_2 (3).docx and applies the May 2025 project guidelines. The source file was preserved. This record accompanies the corrected chapter and is not part of its literature review.')
responses=[
('0','2.1','Replaced the introduction with purpose, funnel structure, themes, critical-comparison criteria and the intended gap; removed implementation detail.'),
('1','2.2 throughout','Compared named studies by task, method, data, findings, limitation, relevance and difference from the proposed study; qualified evidence types.'),
('2','2.2.1','Added survey and academic evidence; distinguished banking methods from wallet and agent-mediated operations.'),
('3','2.2.2','Separated documented misuse from prevalence estimates; distinguished laundering from fraud and other irregularities.'),
('4','2.2.4','Compared graph-learning findings with network-derived tabular summaries; explained the auditability and scope rationale and the loss of graph information.'),
('5','2.2.5; 2.2.14','Added Kenya–Tanzania interviews, African stakeholder research and Zimbabwean agent-policy analysis; identified the absence of classifier validation in those sources.'),
('6','2.2.6; Table 2.1','Corrected the supervised-label contradiction; separated experimental reference labels, model predictions and independent rule alerts.'),
('7','2.2.7','Retained algorithms only when discussing cited researchers; removed the project’s old candidate-selection and final-model assertions.'),
('8','2.2.8','Used leakage literature to state evaluation requirements; reserved the project’s detailed mechanisms for Chapter 3.'),
('9','2.2.9','Compared synthetic-data benefits and validity limits; removed project-specific dataset size and generation detail.'),
('10','2.2.11','Added academic explanation and investigator-support evidence; distinguished model attribution, rule explanation and historical context.'),
('11','2.2.12','Discussed timeliness and measurement boundaries without describing the application’s detailed processing sequence.'),
('12','2.2.13','Strengthened peer-reviewed African evidence and regional AML context; stated the specific regional evidence gap.'),
('13','2.2.14','Made Zimbabwe the final narrowing point; compared legal scholarship and the 2026 banking graph study without claiming no prior local AML research.'),
('14','2.4','Reduced proposed work to two paragraphs explaining the bounded response; retained the separately required feasibility analysis.')]
h(g,'1. Comment-by-Comment Response',2)
table(g,['Comment ID','Revised location','Action'],responses,[.85,1.3,4.15])
h(g,'2. Guideline and Reference Corrections',2)
p(g,'The chapter retains the required Introduction, Analysis of Related Work, Gaps Identified, Proposed Work, Feasibility Analysis and Summary structure. It includes technical, economic, social and operational feasibility, with the missing 2.4.1.1 heading restored. Tables are numbered by chapter and captioned above. The document uses Times New Roman 12-point text, 14-point chapter headings, bold numbered subheadings, 1.5 body spacing, single-spaced captions and Arabic body-page numbers. The supplied student identities were preserved; unsupplied full-dissertation cover details were not invented.')
p(g,'References were checked against publisher, author, university or institutional sources. Both Zimbabwe mobile-money articles were corrected to Howard Chitimira and Elfas Torerai; the 2023 article’s journal was corrected to De Jure. James Whisker’s initial was corrected. The 2015 social-network source was removed to meet the ten-year recency requirement. The revised reference set spans 2019–2026. The ESAAMLG report was distinguished from an effectiveness evaluation.')
p(g,'The chapter uses a complete IEEE-style reference list numbered in order of first citation within this standalone document. When assembling the full dissertation, merge duplicate sources with the Chapter 1 bibliography and renumber citations globally; do not paste two independently numbered lists into one master reference section.')
h(g,'3. Alignment and Evidence Boundaries',2)
p(g,'The proposed work refers to the corrected Chapter 1’s five SMART objectives and 12-week schedule. No old feature counts, dataset dimensions, selected Gradient Boosting model or 0.35 threshold were carried into the literature review. Detailed V4 implementation facts remain matters for Chapters 3 and 4. Those chapters were not edited by this revision.')
p(g,'The chapter does not invent model results, stakeholder studies, measured cost savings or regulatory approval. The US$0 additional procurement assumption covers reuse of existing hardware and the open-source application stack only. Unmeasured utility costs, staff effort, latency, fairness and operator-scale performance are identified as evidence boundaries. This editing task did not retrain models, change the application or execute database transactions.')
g.save(OUT/'Chapter_2_Corrections_Record.docx')
(OUT/'chapter2_reference_map.json').write_text(json.dumps({str(i):key for i,key in enumerate(order,1)},indent=2),encoding='utf-8')
for file in ['Chapter_2_Corrected.docx','Chapter_2_Corrections_Record.docx']:
    dd=Document(OUT/file)
    txt='\n'.join([x.text for x in dd.paragraphs]+[c.text for t in dd.tables for r in t.rows for c in r.cells])
    assert '{cite:' not in txt
    print(file, 'words',len(txt.split()),'tables',len(dd.tables))
assert len(order)==len(refs)
print('All 15 comments mapped; all',len(order),'references cited.')
