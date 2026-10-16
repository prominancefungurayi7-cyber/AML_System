from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent
ROOT = OUT.parent.parent

def base(title, cover=False):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    s.top_margin = s.bottom_margin = Inches(.8)
    s.left_margin, s.right_margin = Inches(1.1), Inches(.85)
    for name in ['Normal', 'Body Text', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
        st = d.styles[name]
        st.font.name = 'Times New Roman'
        st.font.size = Pt(12)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_after = Pt(6)
    d.styles['Normal'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name in ['Heading 1', 'Heading 2', 'Heading 3']:
        d.styles[name].font.bold = True
        d.styles[name].paragraph_format.keep_with_next = True
        d.styles[name].paragraph_format.space_before = Pt(12)
    d.styles['Heading 1'].font.size = Pt(14)
    d.styles['Caption'].paragraph_format.line_spacing = 1
    d.styles['Caption'].font.italic = False
    d.core_properties.title = title
    d.core_properties.author = 'Prominance Fungurayi; Ashton Kahlari'
    d.core_properties.subject = 'Revised Chapter 1 and demonstrable SMART objectives'
    if cover:
        for _ in range(5): d.add_paragraph()
        for text, size, bold in [
            ('MOBILE-MONEY ANTI-MONEY LAUNDERING\nDECISION-SUPPORT SYSTEM',14,True),
            ('CHAPTER 1',14,True),
            ('Prominance Fungurayi: R251946T',12,False),
            ('Ashton Kahlari: R252989W',12,False)]:
            p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(text); r.font.size=Pt(size); r.bold=bold
        s=d.add_section(WD_SECTION_START.NEW_PAGE)
        s.footer.is_linked_to_previous=False
        num=OxmlElement('w:pgNumType'); num.set(qn('w:start'),'1'); s._sectPr.append(num)
    p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); p._p.append(f)
    return d

def h(d, text, level=2): d.add_heading(text, level)
def p(d, text): return d.add_paragraph(text)
def table(d, headers, rows, widths=None):
    t=d.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
    t.autofit=False
    if widths:
        for col,w in zip(t.columns,widths): col.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers):
        c.text=txt
        for r in c.paragraphs[0].runs:r.bold=True
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2');c._tc.get_or_add_tcPr().append(sh)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
    for row in t.rows:
        prop=row._tr.get_or_add_trPr();prop.append(OxmlElement('w:cantSplit'))
        if widths:
            for c,w in zip(row.cells,widths):c.width=Inches(w)
        for c in row.cells:
            for par in c.paragraphs:
                par.alignment=WD_ALIGN_PARAGRAPH.LEFT
                par.paragraph_format.space_after=Pt(3)
                for run in par.runs:run.font.size=Pt(12)
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

objectives=[
('1', 'To implement mobile-money transaction processing by the end of Week 3, with correct balance reconciliation and single processing of repeated requests in all 12 predefined tests covering Cash-In, Cash-Out and Wallet-to-Wallet transactions, including valid payments, invalid requests, insufficient funds and duplicate requests.'),
('2', 'To implement historical behavioural analysis by the end of Week 6 and verify its calculations against independently calculated expected values for six controlled histories: one routine and one suspicious-pattern history for each of structuring, wallet/network and agent behaviour, with no use of future transaction information.'),
('3', 'To evaluate four candidate classification algorithms and integrate one primary classifier by the end of Week 9, documenting a confusion matrix, precision, recall, macro-F1, false-positive rate and training duration for each candidate under the same development-validation protocol, and verifying that stored decisions reproduce the selected model’s outputs.'),
('4', 'To implement an authorised alert-review process by the end of Week 10, passing six predefined tests covering AI-only, rule-only and combined alerts, authorised review and status recording, denial of customer access to compliance cases, and delivery of notifications only to authorised staff.'),
('5', 'To implement reliable post-commit monitoring by the end of Week 12, passing five predefined tests for fresh-history verification, concurrent-history change, inference failure, audited inference retry and repeated payment requests, with no duplicated financial transaction or unexplained balance change in any test.')]

doc=base('Chapter 1: Mobile-Money Anti-Money Laundering Decision-Support System',True)
h(doc,'CHAPTER 1',1)
h(doc,'1.1 Introduction')
p(doc,'Digital financial services have expanded access to payments and transfers through mobile wallets and agent networks. Mobile money has become particularly important in Sub-Saharan Africa, where it supports financial activity beyond conventional bank branches [1]. Zimbabwe’s National Financial Inclusion Strategy II recognised the importance of access to, usage of and quality of financial services [2]. This context provided the motivation for examining transaction-monitoring support for mobile-money services.')
p(doc,'This study addressed the identification of suspicious patterns that may emerge across transaction histories, relationships between wallets and agent activity. It developed a local mobile-money Anti-Money Laundering (AML) decision-support prototype combining behavioural analysis, machine-learning classification and an investigation workflow. Its classifications were intended to support human review rather than establish that money laundering had occurred. This chapter presents the background, problem, aim, objectives, limitations, delimitations, development instruments, work plan and justification.')

h(doc,'1.2 Background of the Study')
p(doc,'International AML monitoring has increasingly considered how technology can assist the assessment of large volumes of financial activity. The Financial Action Task Force (FATF, 2021) examined opportunities for advanced analytics alongside challenges involving data quality, governance and responsible implementation [3]. This institutional review supported investigation of technology-assisted monitoring, but it did not establish that any particular model would be effective in Zimbabwean mobile-money operations. Performance therefore required evaluation within a clearly defined setting.')
p(doc,'Jullum et al. (2020) used supervised machine learning on transaction and investigation data from a Norwegian bank. Their study found that a predictive approach could improve the prioritisation of transactions for investigation compared with the bank’s existing approach [4]. Its relevance was the use of historical information to support AML decisions. However, evidence from one banking environment did not demonstrate equivalent performance for mobile wallets and cash agents; applying its conclusions to Zimbabwe would require further validation.')
p(doc,'Tertychnyi et al. (2022) examined time-aware predictive AML monitoring using customer-history snapshots and an alert policy. Their evaluation considered the timing of detection and the avoidance of redundant alerts, demonstrating the value of assessing a monitoring process beyond a single classification score [5]. The work informed the present study’s attention to chronological behaviour and alert handling. Its institution-specific evaluation, however, did not establish performance for the local mobile-money setting.')
p(doc,'Weber et al. (2019) investigated graph-based financial forensics using Bitcoin transaction data. Their comparison included graph convolutional networks and conventional classifiers, with Random Forest performing strongly in the reported experiments [6]. The study demonstrated a way to examine relationships between transactions and cautioned against assuming that a more complex model necessarily performed better. Cryptocurrency transaction graphs did not represent cash-agent operations directly, but the relational perspective was relevant to wallet-to-wallet monitoring.')
p(doc,'Altman et al. (2023) developed a synthetic financial-transaction generator and AML datasets to support reproducible model comparisons where shareable operational data were scarce [7]. Their contribution provided a controlled basis for investigating transaction patterns. Nevertheless, synthetic benchmarks depended on their simulation assumptions and did not establish effectiveness on confidential operator data. This distinction was relevant to both behavioural evaluation and the interpretation of model results in an academic prototype.')
p(doc,'Within Sub-Saharan Africa, mobile money combined wallet transfers with the activities of cash-in and cash-out agents. The GSMA’s 2025 industry report documented the region’s central role in the mobile-money sector [1]. These service characteristics made the relationship between customer histories, wallet connections and agent activity relevant to monitoring. Regional adoption established the importance of the application area; it did not, by itself, establish the prevalence of money laundering.')
p(doc,'Whisker and Lokanan (2019) reviewed mobile-money characteristics and analysed their implications for AML and counter-terrorist-financing controls. Their thematic analysis identified threats associated with mobile-money services and discussed measures for addressing them [8]. It helped establish why the monitoring of mobile-money channels, including agent-mediated activity, deserved attention. However, it did not provide an experimentally validated agent-behaviour classifier. The present study therefore treated agent activity as a behavioural monitoring dimension whose calculations and predictive usefulness required explicit evaluation.')
p(doc,'In Zimbabwe, Chitimira and Torerai (2021) examined mobile-money regulation, technological innovation and financial inclusion through legal and regulatory analysis. They highlighted the relationship between the regulation of mobile-money providers and inclusive financial services [9]. Their work established local institutional relevance but did not evaluate a machine-learning transaction-monitoring system. The national financial-inclusion strategy also provided a policy context for examining usable and accountable financial services [2].')
p(doc,'Taken together, these studies established the potential of predictive monitoring, historical analysis and transaction relationships, while also identifying data and contextual limitations. The specific gap addressed here was the integration and demonstrable evaluation of those ideas within a local mobile-money prototype that connected transaction integrity, structuring, wallet/network and agent behaviour to an authorised investigation process. The reviewed evidence did not establish how that complete workflow would perform in this setting. The study addressed this bounded implementation and evaluation gap; it did not claim that AML monitoring research or mobile-money controls were absent in Zimbabwe.')

h(doc,'1.3 Problem Definition')
p(doc,'Suspicious mobile-money activity can be distributed across repeated transactions, connected wallets and agent interactions. Reviewing isolated amounts can obscure these relationships, while repetitive alerts can increase investigators’ workloads [4], [5], [8]. The resulting challenge is to assess historical behaviour and prioritise relevant cases without losing transaction context. In the Zimbabwean mobile-money setting, the reviewed evidence did not establish the effectiveness of an integrated local monitoring and investigation workflow. This limits the basis for assessing its practical usefulness.')

h(doc,'1.4 Aim')
p(doc,'To develop a mobile-money AML decision-support system for the Zimbabwean context that integrates transaction processing, historical behavioural analysis, machine-learning classification and authorised investigation of suspicious activity.')
h(doc,'1.5 Objectives')
p(doc,'The study was organised around five process-based objectives within a 12-week development schedule. The measurable criteria below define acceptance requirements; their inclusion does not imply that every test has already passed. Test procedures and achieved results belong in Chapters 3 and 4.')
for num,text in objectives:p(doc,num+'. '+text).paragraph_format.keep_together=True
h(doc,'1.5.1 Algorithm-Specific SMART Commitments',3)
p(doc,'The four candidate-algorithm commitments operationalise Objective 3 and are not additional research objectives. Each required one documented development-validation result using the same data partition and evaluation measures, together with recorded training time in seconds. Completion deadlines were expressed relative to the start of the 12-week schedule. Model selection was reserved for the comparative evaluation, rather than assumed in this chapter.')
p(doc,'Logistic Regression: To establish and evaluate a linear classification baseline by the end of Week 7, producing the confusion matrix and four classification measures specified in Objective 3, together with a recorded training duration for comparison with the other candidates.')
p(doc,'Random Forest: To evaluate an ensemble of decision trees by the end of Week 8, producing the same result set and a recorded training duration under the shared protocol, and comparing its errors with those of the linear baseline.')
p(doc,'Histogram-based Gradient Boosting: To evaluate a boosted-tree candidate by the end of Week 8, producing the same result set and a recorded training duration, and documenting its false-positive and missed-pattern behaviour relative to the other candidates.')
p(doc,'XGBoost: To evaluate a second boosted-tree implementation by the end of Week 9, producing the same result set and a recorded training duration, and completing a four-candidate comparison that supports a documented primary-model selection.')
p(doc,'For these commitments, the four classification measures were precision, recall, macro-F1 and false-positive rate, supported by the confusion matrix; training duration was recorded separately. Measurement of classification quality was required even when a candidate performed poorly. No fixed accuracy percentage or unmeasured execution-speed guarantee was treated as an achieved result.')

h(doc,'1.6 Limitations')
p(doc,'Access to operational data was limited. Confidential mobile-money transaction records and confirmed investigation outcomes were unavailable to the project, so synthetic records were used. Such records allowed controlled testing but reflected the assumptions used to create them and could omit legitimate behaviour, rare suspicious patterns and changes in real customer activity.')
p(doc,'Model evaluation did not establish operational generalisation. Available comparative results were obtained from synthetic development-validation data. They were not an independent external test or a demonstration of effectiveness on a mobile-money operator’s network. Reported classification measures therefore had to be interpreted within that restricted evaluation setting.')
p(doc,'The classifier could miss suspicious patterns and could flag routine activity. Aggregate scores could also conceal differences between structuring, wallet/network and agent scenarios. Historical-data quality and coverage affected the behavioural information available for a decision. These limitations required separate reporting of errors and explicit handling of unavailable inference results.')
p(doc,'Experiments were conducted on one local computer using the hardware described in Section 1.8. Recorded training times reflected that environment and the specific experiment configuration. They did not measure concurrent-user capacity, live transaction latency or operator-scale throughput. The study did not establish how performance would change across other hardware or larger operational workloads; it did not attribute unrecorded experimental restrictions to computing resources.')

h(doc,'1.7 Delimitations')
p(doc,'The study used Zimbabwean mobile-money services as its application context and implemented a local academic prototype. It did not connect to a live operator network, telecommunications infrastructure or external regulatory reporting system. Reports and investigation records produced within the prototype were internal records.')
p(doc,'The transaction scope covered Cash-In, Cash-Out and Wallet-to-Wallet operations. The behavioural scope covered structuring, wallet/transaction-network relationships and agent activity. The classification task distinguished routine from suspicious-pattern behaviour, while independent AML rules remained a separate source of alerts.')
p(doc,'The system supported authorised transaction monitoring and case review. It did not replace customer-identification obligations, investigative judgement or legal determinations. Biometric identification, cryptocurrency analysis, live SIM-registration checks, device intelligence and geolocation-based detection were outside the research scope. The study did not claim that its simulation represented all Zimbabwean providers or regulatory jurisdictions.')

h(doc,'1.8 Development Instruments')
p(doc,'Development instruments supported application programming, relational data storage, behavioural analysis, model evaluation and local testing. The versions below identify the inspected local environment; they do not imply that all software components ran within one Python environment.')
h(doc,'1.8.1 Hardware Instruments',3)
p(doc,'Development was performed on a computer with an Intel Core i5-9400H processor rated at 2.50 GHz and approximately 15.73 GiB of operating-system-visible memory, corresponding to the nominal 16 GB memory specification. The machine hosted the local application, database and model-development tools. These specifications provided context for the recorded experiment timings.')
h(doc,'1.8.2 Software Instruments',3)
doc.add_paragraph('Table 1.1: Inspected software instruments and versions', 'Caption')
table(doc,['Instrument / version','Role'],[
('Windows 11, build 26300.9550','Local development operating system.'),
('Python 3.14.3','Application programming and model execution.'),
('Flask 3.1.3','Web routes, request handling and application services.'),
('Flask-SocketIO 5.6.1','Server-side real-time event communication.'),
('Socket.IO browser client 4.7.5','Browser-side event delivery.'),
('React 18.3.1','Interactive browser components.'),
('MySQL Server 8.0.46','Relational storage and transaction management.'),
('mysql-connector-python 9.7.0','Application connection to MySQL.'),
('scikit-learn 1.9.0','Classification, preprocessing and evaluation utilities.'),
('XGBoost 3.4.1','Boosted-tree experiments and the separate model-runtime environment.'),
('NumPy 2.5.0 (application); 2.4.3 (model runtime)','Numerical arrays in the respective environments.'),
('pandas 3.0.3 (application)','Tabular data manipulation.'),
('SciPy 1.18.0; joblib 1.5.3 (model runtime)','Scientific computation and model serialisation.'),
('Git 2.53.0.windows.2','Source-code version control.'),
('HTML, CSS and JavaScript','Browser structure, presentation and interaction; web languages rather than separately versioned Python packages.')],[2.7,3.6])

h(doc,'1.9 Work Plan')
p(doc,'The retained work plan covered 12 relative development weeks. Table 1.2 identifies activities, deadlines and evidence, while Figure 1.1 shows their planned overlap. The schedule defines development milestones rather than retrospectively certifying completion dates. Documentation ran throughout the project, and final verification assessed the five objectives together.')
doc.add_paragraph('Table 1.2: Twelve-week work plan', 'Caption')
plan=[
('Requirements and design','1–2','Defined scope, workflows and acceptance-test matrix.'),
('Transaction processing','2–3','Objective 1: transaction and balance checks.'),
('Controlled histories and behaviour','3–6','Objective 2: six independently checked history fixtures.'),
('Logistic Regression baseline','7','Baseline metrics and training duration.'),
('Random Forest and histogram boosting','7–8','Comparable candidate metrics and training durations.'),
('XGBoost and model comparison','8–9','Objective 3: four-candidate comparison and primary-model integration.'),
('Alert review and access control','9–10','Objective 4: six alert and authorisation checks.'),
('Reliable post-commit processing','10–12','Objective 5: five integrity and recovery checks.'),
('Documentation and demonstration','1–12','Dissertation, evidence register and final demonstration.')]
table(doc,['Activity','Weeks','Required evidence'],plan,[2.2,.65,3.45])
gantt=[('Requirements',1,2),('Transactions',2,3),('Behaviour',3,6),('LR baseline',7,7),('RF and HGB',7,8),('XGB / comparison',8,9),('Alert review',9,10),('Reliable monitoring',10,12),('Documentation',1,12)]
t=table(doc,['Activity']+[str(x) for x in range(1,13)],[[name]+['■' if start<=w<=end else '' for w in range(1,13)] for name,start,end in gantt],[2.1]+[.35]*12)
for row in t.rows:
    for c in row.cells[1:]:
        for par in c.paragraphs:par.alignment=WD_ALIGN_PARAGRAPH.CENTER
        tcMar=OxmlElement('w:tcMar')
        for side in ['left','right']:
            z=OxmlElement('w:'+side);z.set(qn('w:w'),'0');z.set(qn('w:type'),'dxa');tcMar.append(z)
        c._tc.get_or_add_tcPr().append(tcMar)
for row in t.rows[:-1]:
    for c in row.cells:
        for par in c.paragraphs:par.paragraph_format.keep_with_next=True
doc.add_paragraph('Figure 1.1: Twelve-week project Gantt chart; shaded markers indicate planned activity.', 'Caption')

h(doc,'1.10 Justification / Rationale')
p(doc,'The project was justified by the importance of maintaining useful financial services alongside accountable monitoring. Mobile money’s regional role and Zimbabwe’s financial-inclusion priorities made this a relevant application area [1], [2]. Investigators required a way to examine related activity and distinguish the basis of an alert before deciding whether further review was warranted.')
p(doc,'Historical and relational analysis offered a practical basis for examining activity that an isolated transaction view might overlook. Previous predictive-monitoring studies supported investigating machine learning as an aid to prioritisation, while also showing why evaluation needed to account for errors and the alert workflow [4], [5]. A local prototype provided a bounded setting in which those functions could be demonstrated and questioned.')
p(doc,'The contribution was an integrated, testable mobile-money decision-support workflow connecting transaction integrity, behavioural analysis, classification and authorised investigation. Its academic value lay in evaluating how these processes worked together and documenting their limitations. Its potential practical value was to help investigators trace alerts to their recorded basis while preserving financial and audit records. The project did not claim a new learning algorithm, proven crime reduction or operational effectiveness beyond the evidence collected.')

h(doc,'1.11 Summary')
p(doc,'This chapter established the difficulty of assessing suspicious activity across mobile-money histories and relationships. It defined an overall development aim and five measurable process objectives within a 12-week schedule. It also stated the study’s local scope, data and evaluation limitations, development instruments and justification. The proposed contribution concerned an integrated and demonstrable investigation-support workflow. Chapter 2 reviews the relevant literature and existing systems in greater depth, providing the basis for the design and evaluation choices developed in later chapters.')

h(doc,'1.12 References')
refs=[
'[1] GSMA, The State of the Industry Report on Mobile Money 2025, 2025. [Online]. Available: https://www.gsma.com/sotir/wp-content/uploads/2025/04/The-State-of-the-Industry-Report-2025_English.pdf. Accessed: Oct. 14, 2026.',
'[2] Reserve Bank of Zimbabwe, Zimbabwe National Financial Inclusion Strategy II, 2022–2026, 2022. [Online]. Available: https://www.rbz.co.zw/documents/BLSS/FinancialInclusion/Zimbabwe_National_Financial_Inclusion_Strategy_II_2022-2026.pdf. Accessed: Oct. 14, 2026.',
'[3] Financial Action Task Force, Opportunities and Challenges of New Technologies for AML/CFT, Paris, France, 2021. [Online]. Available: https://www.fatf-gafi.org/content/dam/fatf-gafi/guidance/Opportunities-Challenges-of-New-Technologies-for-AML-CFT.pdf. Accessed: Oct. 14, 2026.',
'[4] M. Jullum, A. Løland, R. B. Huseby, G. Ånonsen, and J. Lorentzen, “Detecting money laundering transactions with machine learning,” Journal of Money Laundering Control, vol. 23, no. 1, pp. 173–186, 2020, doi: 10.1108/JMLC-07-2019-0055.',
'[5] P. Tertychnyi, M. Godgildieva, M. Dumas, and M. Ollikainen, “Time-aware and interpretable predictive monitoring system for Anti-Money Laundering,” Machine Learning with Applications, vol. 8, Art. no. 100306, 2022, doi: 10.1016/j.mlwa.2022.100306.',
'[6] M. Weber et al., “Anti-Money Laundering in Bitcoin: Experimenting with Graph Convolutional Networks for Financial Forensics,” KDD Workshop on Anomaly Detection in Finance, 2019. [Online]. Available: https://arxiv.org/abs/1908.02591. Accessed: Oct. 14, 2026.',
'[7] E. Altman, J. Blanuša, L. von Niederhäusern, B. Egressy, A. Anghel, and K. Atasu, “Realistic Synthetic Financial Transactions for Anti-Money Laundering Models,” Advances in Neural Information Processing Systems, vol. 36, 2023. [Online]. Available: https://arxiv.org/abs/2306.16424. Accessed: Oct. 14, 2026.',
'[8] J. Whisker and M. E. Lokanan, “Anti-money laundering and counter-terrorist financing threats posed by mobile money,” Journal of Money Laundering Control, vol. 22, no. 1, pp. 158–172, 2019, doi: 10.1108/JMLC-10-2017-0061.',
'[9] H. Chitimira and E. Torerai, “The Nexus between Mobile Money Regulation, Innovative Technology and the Promotion of Financial Inclusion in Zimbabwe,” Potchefstroom Electronic Law Journal, vol. 24, 2021, doi: 10.17159/1727-3781/2021/v24i0a10739.'
]
for ref in refs:p(doc,ref).paragraph_format.keep_together=True
doc.save(OUT/'Chapter_1_Corrected.docx')

guide=base('Supervisor Corrections and Panel Demonstration Guide')
h(guide,'SUPERVISOR CORRECTIONS AND PANEL DEMONSTRATION GUIDE',1)
p(guide,'Companion to the revised Chapter 1. This guide records the response to the 12 supervisor comments and explains how the objectives can be demonstrated. It is not an additional dissertation chapter or evidence that the proposed acceptance tests have all passed. The original document and application were left unchanged.')
h(guide,'1. Supervisor Corrections')
corrections=[
('0 / Introduction','1.1','Rewritten opening; concise purpose and chapter map; removed feature counts, dataset composition, model selection, threshold and repeated scope statements.'),
('1 / Background','1.2','Added named studies with methods, findings, limits and relevance; used an international-to-Zimbabwe progression; stated a bounded integration gap and removed project implementation detail.'),
('2 / Problem','1.3','One short problem paragraph; transactions terminology; difficulty, monitoring limitation and consequence; removed the proposed solution.'),
('3 / Aim','1.4','One development verb and one sentence linking behavioural analysis, machine learning and investigation support.'),
('4 / Transaction objective','1.5, Objective 1','Week 3 deadline, three operation types, 12 acceptance cases, balance reconciliation and duplicate-request protection.'),
('5 / Behaviour objective','1.5, Objective 2','Week 6 deadline, six histories across three behavioural dimensions, independently calculated expected values and chronological correctness.'),
('6 / Investigation objective','1.5, Objective 4','Week 10 deadline and six explicit checks for alert provenance, review, recorded status, access and notifications.'),
('7 / Limitations','1.6','Explained synthetic-data necessity and representativeness; limited generalisation; acknowledged model errors; described actual local timing scope without inventing resource restrictions.'),
('8 / Delimitations','1.7','Combined deployment exclusions; complete sentences; maintained internal reports, rules, investigations and supported transaction types.'),
('9 / Versions','1.8','Added inspected hardware and software versions; distinguished application and model-runtime environments.'),
('10 / Justification','1.10','Focused on local value, investigator needs and an integrated academic contribution; removed implementation details and the unsupported reference to real data.'),
('11 / Summary','1.11','Short summary of problem, objectives, scope and contribution; explicit transition to Chapter 2.')]
table(guide,['Comment','Location','Correction'],corrections,[1.2,1.15,3.95])
p(guide,'Additional guideline corrections: five process-based objectives beginning with “To”; dashboard construction replaced by reliable monitoring as the fifth objective; Times New Roman 12-point body text, 14-point chapter heading, numbered headings, 1.5 line spacing, numbered tables and a 12-week Gantt chart. IEEE-style chapter-local references correct the Zimbabwe study’s authors to Chitimira and Torerai. When combining the dissertation, merge these entries into the master bibliography and renumber citations consistently across all chapters.')
h(guide,'2. Demonstrating the Five Objectives')
p(guide,'Use prepared fixtures and evidence from a dedicated demonstration environment. Do not reset the working database or execute live payments merely to complete this guide. The durations below are presentation allocations, not measured algorithm runtimes. Financial demonstrations require the agreed test authorisation and before/after reconciliation. Candidate-model comparison can use the retained experiment records without retraining.')
demos=[
('1 / Week 3 / 4 minutes','Show the 12-case transaction matrix: four cases per operation (valid, invalid, insufficient funds and repeated request). Display expected and actual wallet/agent ledger effects.','All 12 cases match the expected result. Repeated request IDs produce one financial effect. Cash operations require configured demo-agent balances or isolated fixtures.'),
('2 / Week 6 / 4 minutes','Trace six saved histories, one routine and one suspicious-pattern example per behavioural family. Show source records, calculation cut-off and independently calculated expected feature values.','All six calculation comparisons agree within documented numerical tolerance. Later records do not enter earlier feature calculations. This is feature correctness, not a promise that all suspicious examples are detected.'),
('3 / Week 9 / 5 minutes','Open the four-candidate results and timing records. Explain the selection trade-off. Show saved normal and suspicious outputs and their corresponding stored decisions.','Complete, comparable metrics for four candidates; traceable model identity and exact decision agreement. Missing or failed inference is identified explicitly. Report misclassifications honestly.'),
('4 / Week 10 / 4 minutes','Show AI-only, rule-only and combined alerts; review and record a case status; attempt customer access; verify staff-only notification delivery.','All six cases pass. Rule findings remain separate from AI findings, and customer access to compliance data/events is denied.'),
('5 / Week 12 / 4 minutes','Present isolated tests for fresh history, concurrent change, inference failure, audited retry and duplicate payment requests. Display transaction IDs, audit entries and balances.','All five cases preserve financial integrity. Verification or inference failure does not replay a payment, fabricate a prediction or silently select the legacy model.')]
for name,action,pass_rule in demos:
    h(guide,'2.'+name[0]+' Objective '+name,3)
    p(guide,'Demonstration: '+action)
    p(guide,'Acceptance: '+pass_rule)
p(guide,'The 12/6/6/5 case counts are explicit acceptance criteria introduced in this revision. Existing tests should be mapped to these cases; any uncovered case must be marked pending. This document revision did not execute or certify the application tests. Chapters 3 and 4 should contain the case IDs, expected outcomes, actual outcomes and supporting evidence.')
h(guide,'3. SMART Commitments for Each Learning Algorithm')
p(guide,'Specific: evaluate each named candidate for the same suspicious-pattern classification task. Measurable: record the confusion matrix, precision, recall, macro-F1, false-positive rate and training time. Achievable: retained results already exist for all four candidates. Relevant: compare detection errors and computational cost to support the primary-model decision. Time-bound: complete the individual milestones in Weeks 7–9 and the shared comparison by the end of Week 9. These are the retained schedule’s milestones, not a claim that experiment timestamps independently prove the original development start date.')
runtime=json.loads((ROOT/'ml_v4/evidence/stage18/resource_runtime_report.json').read_text())
comparison=json.loads((ROOT/'ml_v4/evidence/stage18/comparison.json').read_text())
rows=[]
for code,name,deadline in [('LR','Logistic Regression','Week 7'),('RF','Random Forest','Week 8'),('HGB','Histogram-based Gradient Boosting','Week 8'),('XGB','XGBoost','Week 9')]:
    v=comparison['primary_'+code]['validation']
    rows.append((name,deadline,f"{runtime['per_fit']['primary_'+code]['fit_seconds']:.3f}",f"{v['positive_precision']:.4f} / {v['positive_recall']:.4f}"))
guide.add_paragraph('Table 3.1: Retained candidate training times and development-validation results','Caption')
table(guide,['Candidate','Deadline','Fit time (s)','Precision / recall'],rows,[2.25,.85,1.05,2.15])
p(guide,'Source: ml_v4/evidence/stage18/resource_runtime_report.json and comparison.json, primary candidate runs. Recorded configuration: one thread, 60,000 development-training rows and 15,000 development-validation rows. Times measure estimator fitting in that experiment, not feature construction, application start-up or payment-to-alert latency. They are observations, not guaranteed maximum runtimes. The full metrics and confusion matrices are retained in comparison.json.')
p(guide,'For the panel, allocate one minute per algorithm to show its saved metrics and observed fit time, then one minute to explain selection. A fresh training run is unnecessary. If end-to-end inference speed is claimed, separately measure elapsed time from committed transaction to persisted inference outcome on the stated hardware and history size. No such latency measurement was made during this document revision.')
h(guide,'4. Alignment with the Current V4 System')
p(guide,'The current frozen authority and retained dissertation evidence identify XGBoost as the primary V4 model, with 40 semantic features, 42 estimator input columns and a 0.50 decision threshold. These implementation facts belong in the methodology and results chapters. The revised introduction and background no longer repeat the old 30-feature architecture, Gradient Boosting selection or 0.35 threshold. The four compared algorithms were Logistic Regression, Random Forest, Histogram-based Gradient Boosting and XGBoost.')
p(guide,'The retained XGBoost development-validation precision was 0.9587 and recall was 0.9036. These figures are synthetic development-validation results, not independent-test or real-world performance. Histogram-based Gradient Boosting had higher recall and macro-F1 in the same comparison, whereas XGBoost had a lower false-positive rate. Explain the documented selection trade-off; do not claim that XGBoost was best on every measure. The retained summary also records uneven detection across scenario families, so overall metrics must not be presented as proof of uniform detection.')
p(guide,'The evidence summary states that the original final-test and independent-evaluation partitions were quarantined. Do not describe their results as successful independent validation. Do not confuse the generator’s routine/suspicious scenario proportions with model accuracy or require predictions to follow the generated class ratio. Independent rule alerts and V4 classifications should remain distinct in the presentation.')
p(guide,'Evidence inspected: ml_v4/evidence/stage23/dissertation_results_summary.md; ml_v4/frozen_inference/authority.json; and the two Stage18 files cited above. These were read, not modified. Chapters 2–5 were not rewritten as part of this Chapter 1 correction; align their model descriptions, evaluation labels, objectives and bibliography before combining the final dissertation.')
guide.save(OUT/'Corrections_and_Panel_Demonstration_Guide.docx')

for path in OUT.glob('*.docx'):
    d=Document(path)
    txt='\n'.join(x.text for x in d.paragraphs)
    print(path.name, 'paragraphs',len(d.paragraphs),'tables',len(d.tables),'words',len(txt.split()))
print('Created corrected chapter and companion guide; source document unchanged.')
