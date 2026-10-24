"""Create the requested printable system report without editing the chapters."""
from pathlib import Path
import json
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
DESTINATION = ROOT / 'Chapters 1-5'


def content():
    recovery = json.loads((ROOT / 'instance/supervisor_feedback_recovery.json').read_text(encoding='utf-8'))
    outcomes = [item.get('result', {}) for item in recovery['actions']]
    if not recovery.get('protected_data_unchanged') or any(item.get('status') != 'success' for item in outcomes):
        raise RuntimeError('Finish and verify the audited recovery before producing the final report.')
    normal = sum(item['decision'] == 'normal' for item in outcomes)
    suspicious = len(outcomes) - normal
    blocks = []
    def h(text): blocks.append(('heading', text))
    def p(text): blocks.append(('paragraph', text))
    def table(rows, widths): blocks.append(('table', (rows, widths)))
    def code(text): blocks.append(('code', text))
    h('1. Purpose and verified operating status')
    p('EpoCash is a mobile-money anti-money laundering decision-support system. It records customer payments, '
      'evaluates transaction patterns using the frozen V4 XGBoost model and independent AML rules, presents '
      'cases to compliance staff, and supports private communication and internal reporting. This report '
      'documents the implemented workflows, supervisor-feedback corrections and verification completed on 24 October 2026.')
    p('The local Flask application was started successfully with V4 selected, the independent rule engine enabled '
      'and the Redis event relay connected. The checks in this report cover the corrected workflows. '
      'They do not represent external FIU integration, model retraining or a new independent model-accuracy evaluation.')
    table([['Component', 'Verified configuration'],
           ['Primary machine-learning model', 'Frozen V4 XGBoost; no Stage14 fallback on failure'],
           ['Independent AML rules', 'ON'],
           ['Persistent data', 'Local MySQL database: aml'],
           ['Realtime delivery', 'Flask-SocketIO with connected Redis relay'],
           ['Application address', 'http://127.0.0.1:5000']], [0.34, 0.66])

    h('2. Users and responsibilities')
    table([['Role', 'Main functions'],
           ['Customer', 'Register and verify an account; view wallet balance and transactions; transfer value; use agent cash services; communicate with customers and compliance staff.'],
           ['Agent', 'Use the service dashboard to process customer cash-in and cash-out, manage electronic float and inspect service history.'],
           ['Compliance officer', 'Review transactions, separate AI and rule findings, investigate alerts, document cases, file SARs and send case documents to administrators.'],
           ['Administrator', 'Manage application users and operational controls, generate demonstration transactions, review reports and receive confidential case documents through Messages.']], [0.25, 0.75])
    p('Role checks protect staff reports and private attachments. Customer-to-administrator messaging remains restricted. '
      'SAR and escalation documents are delivered through the existing application Messages feature, as requested; no telephone SMS service is used.')

    h('3. Transaction processing and financial safety')
    p('A payment is validated against its type, amount, participants, available balance and unique client request ID. '
      'The transaction, corresponding wallet movements, canonical event and independent rule evidence are committed together. '
      'A repeated request with the same identity returns the recorded payment instead of moving the money again.')
    p('After the financial commit, the existing authorised history mechanism verifies the source revision and required '
      'capture records. V4 then builds its causal feature vector and makes a decision. A history or inference error remains '
      'an explicit analytical error; it neither reverses nor replays a committed payment. Subsequent recovery uses the '
      'append-only inference-attempt audit. Original terminal results are retained.')
    p('The existing $15,000 per-transaction bound is retained in the frozen canonical contract. Unsupported amounts are '
      'rejected before commitment. The fixes do not raise this limit or change the model contract.')

    h('4. Cash-in and cash-out corrections')
    table([['Service', 'Customer electronic wallet', 'Agent electronic float', 'Physical cash'],
           ['Cash-in', 'Credited by the service amount', 'Debited by the same amount', 'Customer gives cash to agent'],
           ['Cash-out', 'Debited by the service amount', 'Credited by the same amount', 'Agent gives cash to customer']], [0.15, 0.28, 0.28, 0.29])
    p('The customer form now selects an active agent and derives the canonical service fields automatically. '
      'The customer or agent must confirm the physical cash exchange. Both customer-initiated services and the dedicated '
      'agent route use the same electronic-value transfer logic. Generated cash services also debit and credit the agent correctly.')
    p('The combined customer and agent electronic balance is conserved. An insufficient customer balance or insufficient '
      'agent float rejects the service without a partial balance update. Agent float can be funded by a transfer to its '
      'active wallet or by a customer cash-out. Existing funds are not increased merely to make a demonstration pass.')
    p('At the read-only check, service agent AGENT0001 had $0.00 electronic float. It therefore needs a funded-wallet transfer '
      'or a customer cash-out before it can provide cash-in. Physical cash availability is confirmed by the operator; '
      'the application does not maintain a separate physical-cash inventory ledger.')

    h('5. Registration corrections')
    p('New customer usernames must contain 3 to 60 characters and at least two letters. Letters, numbers, spaces, dots, '
      'hyphens, underscores and apostrophes are permitted. Values such as 0, P and purely numeric names are rejected '
      'before an OTP is sent. Pending registrations are revalidated when their verification code is submitted.')
    p('A successfully verified new customer receives an opening electronic balance of $0.00. Existing account balances '
      'are preserved. Funding subsequently comes through a legitimate transfer or an agent cash-in with sufficient float. '
      'Existing password, authentication and verification controls remain applicable.')

    h('6. V4 analysis and the risk-scoring mechanism')
    p('The primary pipeline uses the frozen V4 XGBoost artifact with 40 semantic input features and 42 preprocessed inputs. '
      'Its transaction-pattern families cover structuring, network activity and agent behaviour. The model artifact, '
      'feature contract and decision threshold of 0.50 have not been changed. Independent rules are evaluated separately.')
    p('For new successful decisions, 70% of the AI score and 30% of the rule score contribute to the overall operational '
      'risk score. The AI component is the model suspicious-probability output on a 0–100 scale. It is not a measured '
      'probability that the model is correct. The rule component represents independent policy indicators. '
      'Review categories may retain an independent rule or AI escalation even when the combined score is lower.')
    p('AI findings and rule findings are displayed separately. Normal AI decisions do not need a suspicious-activity '
      'explanation. Pending and failed reviews are shown explicitly; the system does not invent a prediction or silently use a legacy model.')

    h('7. Transaction generation and performance')
    p('The default demonstration batch is 30 transactions. In mixed mode, 21 are routine scenarios and 9 are suspicious '
      'scenarios: 3 structuring, 3 network and 3 agent scenarios when an active linked agent is available. Larger supported '
      'batches remain selectable. The 70:30 split describes the generated inputs, not a forced model result or alert count.')
    p('Customer selection includes existing registered customers. Routine amounts and counterparties use the available '
      'recent normal or low-risk transaction history. Unfunded wallets can receive value but cannot spend funds they '
      'do not have. Limited history is treated as a limited baseline; long-term agent features retain their original '
      'history-eligibility requirements.')
    p('New agent scenario plans start with cash-out before the corresponding cash-in, allowing an initially empty agent '
      'float to receive genuine electronic value from a funded customer. Older signed plan identities remain unchanged '
      'for receipt and retry verification. A batch plan is computed once and reused for its receipt checks.')
    p('The generation response no longer repeats full history attestation before returning a saved-batch confirmation. '
      'The background V4 coordinator still performs the required attestation before inference. Dense network-path '
      'evaluation now uses exact temporal dynamic programming in the integration adapter. Path counts, amount-spread '
      'ranking and chronological tie-breaks are preserved; frozen feature files are unchanged.')
    p('On existing transaction 21231, read-only inference completed in approximately 5.4 seconds in one run; the full '
      'isolated worker completed in approximately 3.2 seconds in another run. Earlier evaluation of this dense history '
      'ran for several minutes. These are measured single-transaction observations, not a fixed latency guarantee for '
      'every batch. History checks and individual reviews still require time.')

    h('8. SAR retention and administrator handoff')
    p('Filed SAR narratives remain stored in MySQL and can be opened in full from Reports or the corresponding alert. '
      'Reports now include pagination so older SARs remain accessible. Administrators and compliance officers can '
      'download a confidential PDF containing the reference, account, case, author, dates, status and complete narrative.')
    p('Filing a SAR or escalating a case sends a document through a private conversation to administrator accounts. '
      'Recording internal submission sends the updated SAR status. Deterministic attachment identities prevent an '
      'identical retry from delivering the same document twice. A failed delivery leaves the case intact and exposes '
      'a Send to administrator retry action. An administrator who authors a report can access it directly in Reports; '
      'the system does not create an inaccessible self-conversation.')
    p('Attachments use authenticated access and the existing message history, delivery and read-receipt mechanisms. '
      'Customer access to SAR views and confidential report attachments is denied. These statuses represent internal '
      'case handling; no external FIU transmission is claimed.')

    h('9. Verification and historical recovery')
    table([['Check', 'Result'],
           ['Focused supervisor and reporting workflows', '16 checks passed, including registration, cash accounting, insufficient funds, private PDFs, pagination and delivery retries.'],
           ['Browser cash and generation checks', '3 checks passed using the actual frontend code and mocked network requests.'],
           ['Frozen feature equivalence', '49 cases passed against the frozen extractor, including all 40 features, predictions, nonuniform paths and chronology.'],
           ['Audited retry and route protections', '13 checks passed, including source-revision changes during inference.'],
           ['Final plan reuse and 70:30 mix', '1 focused check passed; current local 30-transaction plan also validated without execution.'],
           ['Live local readiness', 'Application health OK; V4 selected; rules ON; Redis connected.']], [0.38, 0.62])
    p(f'The final read-only review found 53 historical AI failures: 46 resource errors and 7 history errors. '
      f'All {len(outcomes)} were recovered through the existing audited nonfinancial retry mechanism: '
      f'{normal} normal decisions and {suspicious} suspicious-pattern decisions. Each retry retained its original '
      'terminal error record and appended attempt evidence. Whole-ledger hashes confirmed unchanged payment fields, '
      'wallet balances, canonical events, capture records and original V4 result records.')
    p('Financial test writes were restricted to disposable test databases. No new live payment was generated during '
      'verification. Live recovery changed analytical decisions, related alert state and append-only retry evidence only. '
      'No database reset, schema alteration, model retraining or frozen-artifact modification was performed.')
    p('The verified source database contained 14,411 existing transactions and 24 users. The authenticated MySQL principal '
      'was aml@127.0.0.1 on database aml, with server UUID d441b7bb-43f6-11f1-a2d8-0a0027000005. '
      'Passwords and connection secrets are excluded from this report.')

    h('10. Suggested panel demonstration')
    for text in [
        '1. Open the local application and identify the customer, agent, compliance and administrator dashboards.',
        '2. Show the registration requirements and a newly verified customer wallet with a zero opening balance.',
        '3. Use a funded demonstration customer and record both starting balances. Demonstrate cash-out first, then cash-in, confirming the physical exchange and reconciling both electronic balances. These are real ledger operations and should be deliberately authorised.',
        '4. As administrator, choose the 30-transaction mixed batch. Show the planned 21 routine and 9 suspicious scenarios before submitting it once. Avoid repeatedly clicking Generate or launching a large batch during the explanation.',
        '5. Follow live review progress. Open a transaction and explain its AI score, rule score and overall score separately. Explain the observed activity from recorded history rather than claiming that a scenario label guarantees a prediction.',
        '6. As compliance officer, document a case and file a SAR. Open the full report and download the PDF. As administrator, open the private document in Messages and demonstrate its delivery/read status.',
        '7. Show the retained transaction history and retry audit to demonstrate that analytical failures do not duplicate payments.'
    ]: p(text)

    h('11. Startup, recovery and operational notes')
    p('From the project directory, activate the project virtual environment and start the existing server. '
      'The local startup bootstrap configures V4, independent rules and the local Redis relay. MySQL must be running.')
    code('.venv\\Scripts\\python.exe server.py')
    code('AML_PRIMARY_MODEL=V4\nAML_RULE_ENGINE_ENABLED=ON\nSOCKETIO_ALLOWED_ORIGINS=http://127.0.0.1:5000')
    p('Open http://127.0.0.1:5000. If a healthy instance is already running, do not start a duplicate process. '
      'After a backend code update, stop that Flask instance and start it once with the updated files. Refresh the '
      'browser to load the updated interface assets.')
    p('For an inference failure, use the staff audited retry action after resolving its cause; never repeat the payment. '
      'For failed SAR handoff, use Send to administrator on the report or case page. Preserve case and attempt evidence. '
      'If deployment rollback is required, stop the app and restore reviewed application code while retaining the database, '
      'uploads and audit records. Do not switch silently to Stage14 or rewrite committed history.')
    p('The remaining demonstration preparation is operational: ensure a funded customer is available, fund agent float '
      'before cash-in, and choose a small batch. Live controlled cash-service and SAR-message demonstrations should '
      'be performed deliberately by their authorised account holders.')
    return blocks


def build():
    blocks = content()
    DESTINATION.mkdir(exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ('Normal', 'Heading 1', 'Heading 2', 'Title'):
        style = document.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != 'Normal'
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.15
    document.styles['Heading 1'].paragraph_format.space_before = Pt(10)
    for title in ('EPOCASH MOBILE-MONEY ANTI-MONEY LAUNDERING SYSTEM', 'SYSTEM REPORT', '24 October 2026'):
        paragraph = document.add_paragraph(title)
        paragraph.alignment = 1
        paragraph.runs[0].bold = True
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('EpoCash System Report | Page ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

    styles = {
        'paragraph': ParagraphStyle('body', fontName='Times-Roman', fontSize=12, leading=15,
                                    spaceAfter=8),
        'heading': ParagraphStyle('heading', fontName='Times-Bold', fontSize=12, leading=15,
                                  spaceBefore=10, spaceAfter=8, keepWithNext=True),
        'code': ParagraphStyle('code', fontName='Courier', fontSize=9, leading=12, spaceAfter=8),
        'title': ParagraphStyle('title', fontName='Times-Bold', fontSize=12, leading=15,
                                spaceAfter=8, alignment=TA_CENTER),
        'cell': ParagraphStyle('cell', fontName='Times-Roman', fontSize=12, leading=15),
    }
    story = [Paragraph(escape(text), styles['title']) for text in
             ('EPOCASH MOBILE-MONEY ANTI-MONEY LAUNDERING SYSTEM', 'SYSTEM REPORT', '24 October 2026')]
    width = A4[0] - 108
    for kind, value in blocks:
        if kind == 'table':
            rows, fractions = value
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.autofit = False
            for row_index, row in enumerate(rows):
                cells = table.add_row().cells
                for i, text in enumerate(row):
                    cells[i].width = Inches(6.77 * fractions[i])
                    paragraph = cells[i].paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.add_run(text).bold = row_index == 0
                properties = table.rows[-1]._tr.get_or_add_trPr()
                properties.append(OxmlElement('w:cantSplit'))
                if row_index == 0:
                    properties.append(OxmlElement('w:tblHeader'))
            properties = table._tbl.tblPr
            spacing = OxmlElement('w:tblCellSpacing'); spacing.set(qn('w:w'), '0'); spacing.set(qn('w:type'), 'dxa'); properties.append(spacing)
            borders = OxmlElement('w:tblBorders')
            for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                element = OxmlElement('w:' + edge)
                for key, val in [('val', 'single'), ('sz', '4'), ('color', '000000'), ('space', '0')]: element.set(qn('w:' + key), val)
                borders.append(element)
            properties.append(borders)
            pdf_rows = [[Paragraph(('<b>' + escape(text) + '</b>') if i == 0 else escape(text), styles['cell'])
                         for text in row] for i, row in enumerate(rows)]
            pdf_table = Table(pdf_rows, colWidths=[width*f for f in fractions], repeatRows=1)
            pdf_table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.5,colors.black),
                ('VALIGN',(0,0),(-1,-1),'TOP'), ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eeeeee')),
                ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
            story.extend([pdf_table, Spacer(1, 8)])
        else:
            paragraph = document.add_paragraph(value, style='Heading 1' if kind == 'heading' else 'Normal')
            if kind == 'code':
                for run in paragraph.runs: run.font.name = 'Consolas'; run.font.size = Pt(10)
            story.append(Paragraph(escape(value).replace('\n', '<br/>'), styles[kind]))
    docx_path = DESTINATION / 'EpoCash System Report.docx'
    document.save(docx_path)
    pdf_path = DESTINATION / 'EpoCash System Report.pdf'
    pdf = SimpleDocTemplate(str(pdf_path), pagesize=A4, leftMargin=54, rightMargin=54,
                           topMargin=48, bottomMargin=48, title='EpoCash System Report', author='EpoCash')
    def footer(canvas, doc):
        canvas.setFont('Times-Roman', 9)
        canvas.drawRightString(A4[0]-54, 25, f'EpoCash System Report | Page {doc.page}')
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    print(json.dumps({'docx': str(docx_path), 'pdf': str(pdf_path), 'blocks': len(blocks)}))


if __name__ == '__main__':
    build()
