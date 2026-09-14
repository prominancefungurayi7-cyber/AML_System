from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


OUT = Path("1_Fungurayi_Prominance_Kahlari_Ashton.pptx")

# StanPro-inspired visual palette
NAVY = RGBColor(9, 30, 66)
BLUE = RGBColor(19, 93, 171)
TEAL = RGBColor(0, 153, 153)
GOLD = RGBColor(243, 183, 44)
RED = RGBColor(202, 62, 71)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(28, 40, 58)
MUTED = RGBColor(91, 108, 128)
PALE = RGBColor(243, 247, 252)
LINE = RGBColor(211, 221, 233)


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.space_after = Pt(0)
    run = p.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, bullets, x, y, w, h, size=19, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(10)
        p.bullet = True
    return box


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line if line else fill
    return shape


def header(slide, title, number, subtitle=None):
    rect(slide, 0, 0, 13.333, 0.22, TEAL)
    add_text(slide, title, 0.62, 0.38, 11.7, 0.52, 28, NAVY, True)
    if subtitle:
        add_text(slide, subtitle, 0.64, 0.94, 11.5, 0.34, 11.5, MUTED)
    add_text(slide, f"{number:02d}", 12.25, 0.37, 0.52, 0.35, 14, TEAL, True, PP_ALIGN.RIGHT)
    rect(slide, 0.62, 7.1, 12.1, 0.01, LINE)
    add_text(slide, "StanPro Bank AML Intelligence Platform  |  Panel 1", 0.62, 7.18, 7.5, 0.18, 8.5, MUTED)


def card(slide, heading, body, x, y, w, h, accent=BLUE):
    rect(slide, x, y, w, h, WHITE, True, LINE)
    rect(slide, x, y, 0.09, h, accent)
    add_text(slide, heading, x + 0.28, y + 0.22, w - 0.45, 0.35, 15, NAVY, True)
    add_text(slide, body, x + 0.28, y + 0.68, w - 0.45, h - 0.82, 13, INK)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# 1. Title
s = prs.slides.add_slide(blank)
rect(s, 0, 0, 13.333, 7.5, NAVY)
rect(s, 0, 0, 13.333, 0.16, TEAL)
add_text(s, "StanPro Bank", 0.78, 0.70, 3.3, 0.4, 16, GOLD, True)
add_text(s, "AML Intelligence Platform", 0.75, 1.25, 8.3, 0.75, 34, WHITE, True)
add_text(s, "A proposed explainable, real-time system for detecting and managing suspicious financial activity", 0.78, 2.12, 8.7, 0.65, 19, RGBColor(210, 224, 242))
rect(s, 0.78, 3.15, 5.65, 1.45, RGBColor(18, 53, 101), True)
add_text(s, "Research Progress Presentation", 1.05, 3.48, 5.1, 0.3, 16, WHITE, True, PP_ALIGN.CENTER)
add_text(s, "Panel 1  •  7 September 2026", 1.05, 3.92, 5.1, 0.28, 13, RGBColor(210, 224, 242), False, PP_ALIGN.CENTER)
add_text(s, "Prominance Fungurayi  |  R251946T\nAshton Kahlari  |  R252989W", 0.8, 5.65, 5.8, 0.63, 16, WHITE)
# visual motif
for x, y, c, label in [(9.15, 1.25, TEAL, "Monitor"), (10.65, 2.45, GOLD, "Detect"), (9.15, 3.7, RED, "Alert")]:
    shape = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(1.25), Inches(1.25))
    shape.fill.solid(); shape.fill.fore_color.rgb = c; shape.line.color.rgb = c
    add_text(s, label, x, y + 0.46, 1.25, 0.25, 12, NAVY, True, PP_ALIGN.CENTER)
add_text(s, "Risk-based • Auditable • Human-led", 8.1, 5.85, 4.2, 0.3, 14, RGBColor(210, 224, 242), True, PP_ALIGN.CENTER)

# 2 Background
s = prs.slides.add_slide(blank); header(s, "Background and Context", 2, "Why financial institutions need faster, explainable AML monitoring")
add_text(s, "Financial crime monitoring must keep pace with rapid, connected transactions.", 0.7, 1.5, 7.7, 0.55, 22, NAVY, True)
card(s, "Operational challenge", "Manual review and disconnected records can delay the identification, investigation and reporting of unusual activity.", 0.72, 2.35, 3.75, 2.55, BLUE)
card(s, "Regulatory expectation", "FATF standards require risk-based customer due diligence, payment transparency, suspicious-transaction reporting and an effective FIU function.", 4.82, 2.35, 3.75, 2.55, TEAL)
card(s, "Technology opportunity", "Rules, behavioural analytics and machine learning can prioritise cases while preserving evidence for human compliance decisions.", 8.92, 2.35, 3.75, 2.55, GOLD)
add_text(s, "Project setting: a prototype bank AML environment aligned to Zimbabwean FIU reporting workflows and international FATF guidance.", 0.72, 5.52, 11.85, 0.36, 15, MUTED)

# 3 Problem
s = prs.slides.add_slide(blank); header(s, "Problem Statement", 3)
rect(s, 0.75, 1.48, 11.82, 1.0, PALE, True, LINE)
add_text(s, "How can a bank detect potentially suspicious transactions promptly, prioritise explainable alerts, and support compliance staff with auditable case and reporting workflows?", 1.08, 1.72, 11.15, 0.48, 21, NAVY, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_bullets(s, [
    "Static thresholds alone may miss structuring, rapid movement of funds and unusual customer behaviour.",
    "High alert volumes can overwhelm analysts when alerts lack clear evidence and prioritisation.",
    "Fragmented transaction, screening and reporting processes reduce traceability and slow review."
], 1.0, 3.05, 11.2, 2.5, 19)
add_text(s, "Research focus: support—not replace—compliance judgement with transparent technology.", 0.98, 5.95, 11.2, 0.34, 16, TEAL, True, PP_ALIGN.CENTER)

# 4 Solution
s = prs.slides.add_slide(blank); header(s, "Proposed Solution", 4, "An integrated AML intelligence platform for transaction monitoring and case management")
steps = [("1", "Ingest", "Transactions and customer data"), ("2", "Screen", "Watchlist / PEP checks"), ("3", "Score", "Rules + behaviour + ML"), ("4", "Investigate", "Alerts, evidence and cases"), ("5", "Report", "CTR / SAR workflow")]
for i, (num, name, desc) in enumerate(steps):
    x = 0.55 + i * 2.55
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.05), Inches(2.1), Inches(2.2))
    shape.fill.solid(); shape.fill.fore_color.rgb = PALE; shape.line.color.rgb = LINE
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.75), Inches(2.30), Inches(0.58), Inches(0.58))
    circ.fill.solid(); circ.fill.fore_color.rgb = [BLUE, TEAL, GOLD, RED, NAVY][i]; circ.line.color.rgb = circ.fill.fore_color.rgb
    add_text(s, num, x + 0.75, 2.42, 0.58, 0.18, 12, WHITE, True, PP_ALIGN.CENTER)
    add_text(s, name, x + 0.15, 3.05, 1.8, 0.28, 16, NAVY, True, PP_ALIGN.CENTER)
    add_text(s, desc, x + 0.18, 3.47, 1.75, 0.45, 11.5, INK, False, PP_ALIGN.CENTER)
    if i < 4:
        add_text(s, "→", x + 2.13, 2.96, 0.3, 0.32, 21, TEAL, True, PP_ALIGN.CENTER)
add_text(s, "Design principle: every alert should retain its reason, rules triggered, risk score and reviewer actions.", 0.9, 5.30, 11.6, 0.45, 17, NAVY, True, PP_ALIGN.CENTER)

# 5 Objectives
s = prs.slides.add_slide(blank); header(s, "Aim and Objectives", 5)
rect(s, 0.72, 1.35, 11.9, 0.75, RGBColor(230, 243, 249), True, RGBColor(195, 225, 235))
add_text(s, "Aim: To design and evaluate a prototype AML intelligence platform that supports early, explainable detection and compliant handling of suspicious transactions.", 1.0, 1.56, 11.35, 0.32, 16.5, NAVY, True, PP_ALIGN.CENTER)
objectives = [
    ("01", "Identify", "Review AML monitoring requirements, typologies and reporting expectations."),
    ("02", "Design", "Define an auditable architecture combining rules, screening and behavioural risk signals."),
    ("03", "Develop", "Implement a prototype with dashboards, alert/case workflows, CTR and SAR support."),
    ("04", "Evaluate", "Test detection behaviour, evidence capture, usability and reporting workflow outcomes.")
]
for i, (n, title, body) in enumerate(objectives):
    y = 2.52 + (i % 2) * 1.72; x = 0.78 + (i // 2) * 6.08
    rect(s, x, y, 5.7, 1.32, WHITE, True, LINE)
    add_text(s, n, x + 0.25, y + 0.37, 0.5, 0.25, 16, TEAL, True)
    add_text(s, title, x + 0.95, y + 0.23, 1.5, 0.3, 16, NAVY, True)
    add_text(s, body, x + 0.95, y + 0.62, 4.35, 0.45, 12.5, INK)

# 6 Lit
s = prs.slides.add_slide(blank); header(s, "Literature Review: Key Findings and Gaps", 6)
add_text(s, "Literature and standards support risk-based, technology-enabled monitoring—but reveal practical implementation gaps.", 0.72, 1.38, 12, 0.36, 17, NAVY, True)
card(s, "What is established", "FATF promotes risk-based AML/CFT controls, timely suspicious-transaction reporting, FIU analysis and payment transparency.", 0.72, 2.05, 3.75, 2.7, BLUE)
card(s, "What technology adds", "Analytics can improve speed, prioritisation and detection of complex patterns; it must also address governance, privacy and explainability.", 4.80, 2.05, 3.75, 2.7, TEAL)
card(s, "Gap addressed by this study", "A locally adaptable prototype that combines rule evidence, customer behaviour, screening, analyst case management and reporting support in one auditable workflow.", 8.88, 2.05, 3.75, 2.7, GOLD)
add_text(s, "Research gap is an inference from the cited standards and technology literature; it will be refined through the full literature review and stakeholder feedback.", 0.75, 5.35, 11.8, 0.36, 12.5, MUTED, False, PP_ALIGN.CENTER)
add_text(s, "Sources: FATF Recommendations (2026); FATF New Technologies report (2021); Deprez et al. (2024).", 0.75, 5.86, 11.8, 0.25, 11, MUTED, False, PP_ALIGN.CENTER)

# 7 approach
s = prs.slides.add_slide(blank); header(s, "Proposed Research and Development Approach", 7)
phases = [
    ("Discover", "Requirements, literature review and AML typology analysis", BLUE),
    ("Design", "Data model, security, roles and risk-scoring architecture", TEAL),
    ("Build", "Flask prototype, dashboards, rules, screening and workflow", GOLD),
    ("Evaluate", "Scenario-based tests, metrics, user feedback and refinement", RED),
]
for i, (name, desc, colour) in enumerate(phases):
    x = 0.7 + i * 3.12
    rect(s, x, 1.85, 2.65, 3.0, WHITE, True, LINE)
    rect(s, x, 1.85, 2.65, 0.13, colour)
    add_text(s, f"0{i+1}", x + 0.25, 2.25, 0.5, 0.25, 15, colour, True)
    add_text(s, name, x + 0.25, 2.78, 2.05, 0.33, 18, NAVY, True)
    add_text(s, desc, x + 0.25, 3.45, 2.08, 0.75, 13, INK)
    if i < 3:
        add_text(s, "→", x + 2.69, 3.15, 0.35, 0.3, 22, TEAL, True, PP_ALIGN.CENTER)
add_text(s, "Evaluation measures: detection of prepared AML scenarios, completeness of evidence, reporting readiness, response time and analyst usability.", 0.78, 5.55, 11.8, 0.4, 15, NAVY, True, PP_ALIGN.CENTER)

# 8 System model
s = prs.slides.add_slide(blank); header(s, "Proposed System Architecture", 8)
for title, body, x, colour in [
    ("Data Layer", "Users\nTransactions\nWatchlist", 0.78, BLUE),
    ("Detection Layer", "AML rules\nBehavioural profiling\nML risk prediction", 4.95, TEAL),
    ("Compliance Layer", "Alerts & cases\nCTR / SAR reports\nAudit log", 9.12, GOLD),
]:
    rect(s, x, 2.0, 3.35, 2.65, WHITE, True, LINE)
    rect(s, x, 2.0, 3.35, 0.16, colour)
    add_text(s, title, x + 0.28, 2.35, 2.75, 0.3, 18, NAVY, True, PP_ALIGN.CENTER)
    add_text(s, body, x + 0.4, 2.95, 2.52, 1.05, 14, INK, False, PP_ALIGN.CENTER)
add_text(s, "→", 4.26, 3.15, 0.42, 0.35, 27, TEAL, True, PP_ALIGN.CENTER)
add_text(s, "→", 8.44, 3.15, 0.42, 0.35, 27, TEAL, True, PP_ALIGN.CENTER)
rect(s, 2.2, 5.3, 8.95, 0.7, RGBColor(230, 243, 249), True, RGBColor(195, 225, 235))
add_text(s, "Cross-cutting controls: role-based access • evidence retention • audit trails • real-time updates • human review", 2.45, 5.53, 8.42, 0.22, 14.5, NAVY, True, PP_ALIGN.CENTER)

# 9 progress
s = prs.slides.add_slide(blank); header(s, "Progress to Date and Next Steps", 9)
add_text(s, "Current progress", 0.78, 1.48, 5.5, 0.33, 19, NAVY, True)
add_bullets(s, [
    "Defined the problem domain, solution concept and core AML workflows.",
    "Designed a working prototype structure for users, transactions, risk scores, alerts, cases, SARs, CTRs and activity logging.",
    "Implemented draft detection components: configurable rule assessment, screening hooks, behavioural profiling and ML risk prediction.",
    "Prepared synthetic transaction scenarios for normal, suspicious and high-risk activity."
], 0.85, 2.05, 5.55, 3.6, 15.5)
add_text(s, "Next steps", 7.0, 1.48, 5.5, 0.33, 19, NAVY, True)
for i, text in enumerate(["Complete targeted literature review and requirements validation.", "Test the prototype with predefined AML scenarios and document results.", "Refine explainability, security controls and reporting outputs.", "Obtain supervisor feedback and maintain the signed progress log book."]):
    y = 2.05 + i * 0.83
    shape = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.06), Inches(y), Inches(0.37), Inches(0.37))
    shape.fill.solid(); shape.fill.fore_color.rgb = TEAL; shape.line.color.rgb = TEAL
    add_text(s, str(i + 1), 7.06, y + 0.08, 0.37, 0.13, 9, WHITE, True, PP_ALIGN.CENTER)
    add_text(s, text, 7.68, y + 0.02, 4.8, 0.44, 15, INK)

# 10 References
s = prs.slides.add_slide(blank); header(s, "References and Closing", 10)
add_text(s, "Selected references", 0.78, 1.42, 4, 0.35, 20, NAVY, True)
refs = [
    "Financial Action Task Force (FATF). (2026). The FATF Recommendations. https://www.fatf-gafi.org/en/publications/fatfrecommendations/fatf-recommendations.html",
    "FATF. (2021). Opportunities and Challenges of New Technologies for AML/CFT. https://www.fatf-gafi.org/en/publications/digitaltransformation/digital-transformation.html",
    "Reserve Bank of Zimbabwe. (2018). Money Laundering and Proceeds of Crime Amendment Act, No. 12 of 2018. https://www.rbz.co.zw/",
    "Deprez, B., et al. (2024). Network Analytics for Anti-Money Laundering: A Systematic Literature Review and Experimental Evaluation. arXiv:2405.19383."
]
add_bullets(s, refs, 0.82, 1.92, 11.7, 3.25, 13, INK)
rect(s, 0.78, 5.55, 11.75, 0.68, NAVY, True)
add_text(s, "Thank you — we welcome your feedback and questions.", 1.02, 5.76, 11.25, 0.24, 18, WHITE, True, PP_ALIGN.CENTER)
add_text(s, "Prominance Fungurayi (R251946T)  •  Ashton Kahlari (R252989W)", 1.0, 6.48, 11.3, 0.26, 13, MUTED, False, PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT.resolve())
